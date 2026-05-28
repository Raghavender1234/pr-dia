import csv
import io
import os
from zoneinfo import ZoneInfo

from flask import Blueprint, jsonify, request, url_for
from app import db
from app.models.class_management import Class, Notification
from app.models.organization import OrganizationMembership
from app.models.user import User
from datetime import date, datetime, time, timedelta, timezone
from .decorators import admin_role_required, org_member_required
from flask_jwt_extended import jwt_required, get_jwt_identity
from app.utils.membership import can_access_internal_org_features

class_bp = Blueprint('class_api', __name__, url_prefix='/api/classes')

def _current_user():
    """Helper function to get the user object from the JWT token."""
    user_id = get_jwt_identity()
    return User.query.get(user_id)

def ensure_utc_aware(dt):
    """
    Ensure a datetime is UTC-aware.
    If naive, assume it's already in UTC and add timezone info.
    If aware, convert to UTC.
    """
    if dt is None:
        return None
    
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    else:
        return dt.astimezone(timezone.utc)

def parse_datetime_to_naive_utc(value, timezone_name='UTC'):
    if value is None or value == '':
        raise ValueError('scheduled_at is required')

    if isinstance(value, datetime):
        parsed = value
    else:
        scheduled_at_str = str(value).strip()
        if scheduled_at_str.endswith('Z'):
            scheduled_at_str = scheduled_at_str[:-1] + '+00:00'
        parsed = datetime.fromisoformat(scheduled_at_str)

    if parsed.tzinfo is None:
        try:
            parsed = parsed.replace(tzinfo=ZoneInfo(timezone_name or 'UTC'))
        except Exception:
            parsed = parsed.replace(tzinfo=timezone.utc)
    else:
        parsed = parsed.astimezone(timezone.utc)

    return parsed.astimezone(timezone.utc).replace(tzinfo=None)

def schedule_class_reminder(class_id, scheduled_at_naive_utc):
    try:
        from app.extensions import scheduler
        from app.tasks import send_class_reminders

        start_time_utc = ensure_utc_aware(scheduled_at_naive_utc)
        reminder_time = start_time_utc - timedelta(minutes=15)
        now_in_utc = datetime.now(timezone.utc)
        run_at = now_in_utc + timedelta(seconds=5) if reminder_time <= now_in_utc else reminder_time

        scheduler.add_job(
            id=f'class_reminder_{class_id}',
            func=send_class_reminders,
            args=[class_id],
            trigger='date',
            run_date=run_at,
            replace_existing=True
        )
    except Exception as scheduler_error:
        print(f"Warning: Could not schedule reminder: {scheduler_error}")

def iter_class_start_times(start_time_naive_utc, repeats='none', repeat_until=None, timezone_name='UTC'):
    repeats = (repeats or 'none').strip().lower()
    if repeats not in {'none', 'daily', 'weekly'}:
        raise ValueError('repeats must be none, daily, or weekly')

    yield start_time_naive_utc

    if repeats == 'none':
        return

    if not repeat_until:
        raise ValueError('repeat_until is required when repeats is daily or weekly')

    repeat_until_value = repeat_until
    if isinstance(repeat_until, date) and not isinstance(repeat_until, datetime):
        repeat_until_value = datetime.combine(repeat_until, time(23, 59, 59))
    elif isinstance(repeat_until, str) and 'T' not in repeat_until and ' ' not in repeat_until:
        repeat_until_value = f'{repeat_until}T23:59:59'

    repeat_until_naive_utc = parse_datetime_to_naive_utc(repeat_until_value, timezone_name)
    step = timedelta(days=1 if repeats == 'daily' else 7)
    current = start_time_naive_utc + step
    count = 1

    while current <= repeat_until_naive_utc:
        if count >= 366:
            raise ValueError('Recurring imports are limited to 366 class instances')
        yield current
        current += step
        count += 1

def create_class_instances(data, current_user, organization):
    timezone_name = data.get('timezone') or 'UTC'
    start_time_naive_utc = parse_datetime_to_naive_utc(data.get('scheduled_at'), timezone_name)
    duration_minutes = int(data.get('duration_minutes') or 60)
    if duration_minutes <= 0:
        raise ValueError('duration_minutes must be greater than zero')

    created_classes = []
    for scheduled_at in iter_class_start_times(
        start_time_naive_utc,
        data.get('repeats', 'none'),
        data.get('repeat_until'),
        timezone_name
    ):
        new_class = Class(
            organization_id=organization.id,
            title=data['title'],
            description=data.get('description', ''),
            instructor_id=current_user.id,
            scheduled_at=scheduled_at,
            duration_minutes=duration_minutes,
            timezone=timezone_name,
            meeting_link=data.get('meeting_link'),
            created_by=current_user.id
        )
        db.session.add(new_class)
        db.session.flush()
        created_classes.append(new_class)

    return created_classes

@class_bp.route('/', methods=['POST'])
@admin_role_required(allowed_roles=['main_co_admin', 'co_admin'])
def create_class(current_user, organization, **kwargs):
    data = request.get_json()
    
    required_fields = ['title', 'scheduled_at']
    if not all(field in data for field in required_fields):
        return jsonify({'error': 'Missing required fields'}), 400

    try:
        created_classes = create_class_instances(data, current_user, organization)
        db.session.commit()

        for new_class in created_classes:
            schedule_class_reminder(new_class.id, new_class.scheduled_at)

        return jsonify({
            'message': f'{len(created_classes)} class(es) created successfully!',
            'id': created_classes[0].id if created_classes else None,
            'ids': [cls.id for cls in created_classes]
        }), 201

    except Exception as e:
        db.session.rollback()
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Error: {str(e)}'}), 500

def normalize_header(header):
    return str(header or '').strip().lower().replace(' ', '_').replace('-', '_')

def combine_date_and_time(date_value, time_value):
    if isinstance(date_value, datetime):
        parsed_date = date_value.date()
    elif isinstance(date_value, date):
        parsed_date = date_value
    else:
        parsed_date = datetime.fromisoformat(str(date_value).strip()).date()

    if isinstance(time_value, datetime):
        parsed_time = time_value.time()
    elif isinstance(time_value, time):
        parsed_time = time_value
    else:
        parsed_time = time.fromisoformat(str(time_value).strip())

    return datetime.combine(parsed_date, parsed_time)

def read_import_rows(upload):
    filename = upload.filename or ''
    ext = os.path.splitext(filename)[1].lower()

    if ext == '.csv':
        text = upload.stream.read().decode('utf-8-sig')
        return list(csv.DictReader(io.StringIO(text)))

    if ext in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook

        upload.stream.seek(0)
        workbook = load_workbook(upload.stream, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [normalize_header(cell) for cell in next(rows, [])]
        return [
            {headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i]}
            for row in rows
            if any(value not in (None, '') for value in row)
        ]

    raise ValueError('Please upload a CSV or XLSX file')

@class_bp.route('/import', methods=['POST'])
@admin_role_required(allowed_roles=['main_co_admin', 'co_admin'])
def import_classes(current_user, organization, **kwargs):
    upload = request.files.get('file')
    if not upload or not upload.filename:
        return jsonify({'message': 'Please choose a CSV or XLSX file.'}), 400

    try:
        rows = read_import_rows(upload)
    except Exception as exc:
        return jsonify({'message': str(exc)}), 400

    created_classes = []
    errors = []

    for row_number, raw_row in enumerate(rows, start=2):
        row = {normalize_header(key): value for key, value in raw_row.items()}
        try:
            title = str(row.get('title') or '').strip()
            if not title:
                raise ValueError('title is required')

            if not row.get('scheduled_at'):
                date_value = row.get('date')
                time_value = row.get('time')
                if date_value and time_value:
                    row['scheduled_at'] = combine_date_and_time(date_value, time_value)

            row['title'] = title
            created_classes.extend(create_class_instances(row, current_user, organization))
        except Exception as exc:
            errors.append({'row': row_number, 'error': str(exc)})

    if not created_classes:
        db.session.rollback()
        return jsonify({'message': 'No classes were imported.', 'errors': errors}), 400

    db.session.commit()

    for new_class in created_classes:
        schedule_class_reminder(new_class.id, new_class.scheduled_at)

    return jsonify({
        'message': f'Imported {len(created_classes)} class(es).',
        'imported_count': len(created_classes),
        'errors': errors
    }), 201


@class_bp.route('/organization/<int:organization_id>', methods=['GET'])
@jwt_required()
def get_classes_for_organization(organization_id):
    user = _current_user()
    membership = OrganizationMembership.query.filter_by(
        user_id=user.id,
        organization_id=organization_id,
        is_active=True
    ).first()
    
    if not membership:
        return jsonify({'error': 'Access denied'}), 403
    if not can_access_internal_org_features(user, membership):
        return jsonify({'error': 'Class schedules are available only to internally invited students.'}), 403

    classes = Class.query.filter_by(organization_id=organization_id).all()

    calendar_events = []
    for cls in classes:
        if not cls.scheduled_at or not cls.duration_minutes:
            continue

        # Assume database stores naive UTC, add timezone info
        start_time_utc = ensure_utc_aware(cls.scheduled_at)
        end_time_utc = start_time_utc + timedelta(minutes=cls.duration_minutes)
        
        # Format as ISO 8601 with Z suffix (UTC indicator)
        start_iso = start_time_utc.strftime('%Y-%m-%dT%H:%M:%S.000Z')
        end_iso = end_time_utc.strftime('%Y-%m-%dT%H:%M:%S.000Z')

        join_url = url_for('frontend.live_class_view', class_id=cls.id, _external=True)

        calendar_events.append({
            'id': cls.id,
            'title': cls.title,
            'start': start_iso,
            'end': end_iso,
            'extendedProps': {
                'description': cls.description or '',
                'meeting_link': cls.meeting_link or '',
                'class_url': join_url
            }
        })
    
    return jsonify(calendar_events)
